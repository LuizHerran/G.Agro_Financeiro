from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from app.utils.dates import database_date_to_brazil


def generate_history_excel(
    transactions,
    summary,
    output_path,
    filters,
):
    output_path = Path(output_path)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Histórico"

    # =========================================================
    # ESTILOS
    # =========================================================

    title_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    section_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="5B9BD5",
    )

    total_fill = PatternFill(
        fill_type="solid",
        fgColor="E2F0D9",
    )

    cancelled_fill = PatternFill(
        fill_type="solid",
        fgColor="FCE4D6",
    )

    border = Border(
        left=Side(style="thin", color="D9E1F2"),
        right=Side(style="thin", color="D9E1F2"),
        top=Side(style="thin", color="D9E1F2"),
        bottom=Side(style="thin", color="D9E1F2"),
    )

    # =========================================================
    # TÍTULO - LINHA 1
    # =========================================================

    worksheet.merge_cells("A1:G1")

    worksheet["A1"] = "CONTROLE FINANCEIRO"

    worksheet["A1"].font = Font(
        size=18,
        bold=True,
        color="FFFFFF",
    )

    worksheet["A1"].fill = title_fill

    worksheet["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    worksheet.row_dimensions[1].height = 30

    # =========================================================
    # PERÍODO - LINHA 2
    # =========================================================

    worksheet.merge_cells("A2:G2")

    period_text = "Período: "

    if filters.get("start_date"):
        period_text += database_date_to_brazil(
            filters["start_date"]
        )
    else:
        period_text += "Início"

    period_text += " até "

    if filters.get("end_date"):
        period_text += database_date_to_brazil(
            filters["end_date"]
        )
    else:
        period_text += "Fim"

    worksheet["A2"] = period_text

    worksheet["A2"].font = Font(
        size=11,
        bold=True,
        color="FFFFFF",
    )

    worksheet["A2"].fill = title_fill

    worksheet["A2"].alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    # =========================================================
    # TRANSAÇÕES - LINHA 3
    # =========================================================

    worksheet.merge_cells("A3:G3")

    worksheet["A3"] = "TRANSAÇÕES"

    worksheet["A3"].font = Font(
        size=12,
        bold=True,
    )

    worksheet["A3"].fill = section_fill

    worksheet["A3"].alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    # =========================================================
    # CABEÇALHO DAS COLUNAS - LINHA 4
    # =========================================================

    headers = [
        "ID",
        "Tipo",
        "Data",
        "Descrição",
        "Valor",
        "% do total",
        "Status",
    ]

    header_row = 4

    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(
            row=header_row,
            column=column,
        )

        cell.value = header

        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )

        cell.fill = header_fill

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        cell.border = border

    # =========================================================
    # VALORES DO RESUMO
    # =========================================================

    total_ganhos = summary["total_ganhos"]
    total_gastos = summary["total_gastos"]

    # Fluxo total = tudo que foi movimentado pelas
    # transações ATIVAS.
    fluxo_total = total_ganhos + total_gastos

    # =========================================================
    # TRANSAÇÕES
    # =========================================================

    first_data_row = header_row + 1

    for row_index, transaction in enumerate(
        transactions,
        start=first_data_row,
    ):
        transaction_type = transaction["tipo"]
        status = transaction["status"]
        value_centavos = transaction["valor_centavos"]

        # ID
        worksheet.cell(
            row=row_index,
            column=1,
            value=transaction["id"],
        )

        # Tipo
        worksheet.cell(
            row=row_index,
            column=2,
            value=transaction_type,
        )

        # Data
        worksheet.cell(
            row=row_index,
            column=3,
            value=database_date_to_brazil(
                transaction["data_transacao"]
            ),
        )

        # Descrição
        worksheet.cell(
            row=row_index,
            column=4,
            value=transaction["descricao"],
        )

        # Valor
        worksheet.cell(
            row=row_index,
            column=5,
            value=value_centavos / 100,
        )

        worksheet.cell(
            row=row_index,
            column=5,
        ).number_format = 'R$ #,##0.00'

        # =====================================================
        # PERCENTUAL
        # =====================================================

        percentage_cell = worksheet.cell(
            row=row_index,
            column=6,
        )

        if status == "CANCELADA":
            percentage_cell.value = "—"

        elif transaction_type == "GANHO":
            if total_ganhos > 0:
                percentage_cell.value = (
                    value_centavos / total_ganhos
                )
            else:
                percentage_cell.value = 0

            percentage_cell.number_format = "0.00%"

        elif transaction_type == "GASTO":
            if total_gastos > 0:
                percentage_cell.value = (
                    value_centavos / total_gastos
                )
            else:
                percentage_cell.value = 0

            percentage_cell.number_format = "0.00%"

        # Status
        worksheet.cell(
            row=row_index,
            column=7,
            value=status,
        )

        # =====================================================
        # FORMATAÇÃO DA LINHA
        # =====================================================

        for column in range(1, 8):
            cell = worksheet.cell(
                row=row_index,
                column=column,
            )

            cell.border = border

            cell.alignment = Alignment(
                vertical="center",
            )

        # Alinhamentos específicos

        worksheet.cell(
            row=row_index,
            column=1,
        ).alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        worksheet.cell(
            row=row_index,
            column=2,
        ).alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        worksheet.cell(
            row=row_index,
            column=3,
        ).alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        worksheet.cell(
            row=row_index,
            column=5,
        ).alignment = Alignment(
            horizontal="right",
            vertical="center",
        )

        worksheet.cell(
            row=row_index,
            column=6,
        ).alignment = Alignment(
            horizontal="right",
            vertical="center",
        )

        worksheet.cell(
            row=row_index,
            column=7,
        ).alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        # =====================================================
        # DESTACA CANCELADAS
        # =====================================================

        if status == "CANCELADA":
            for column in range(1, 8):
                worksheet.cell(
                    row=row_index,
                    column=column,
                ).fill = cancelled_fill

    # =========================================================
    # RESUMO
    # =========================================================

    resumo_row = first_data_row + len(transactions)

    worksheet.merge_cells(
        start_row=resumo_row,
        start_column=1,
        end_row=resumo_row,
        end_column=7,
    )

    worksheet.cell(
        row=resumo_row,
        column=1,
        value="RESUMO",
    )

    worksheet.cell(
        row=resumo_row,
        column=1,
    ).font = Font(
        bold=True,
        size=12,
    )

    worksheet.cell(
        row=resumo_row,
        column=1,
    ).fill = section_fill

    worksheet.cell(
        row=resumo_row,
        column=1,
    ).alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    # =========================================================
    # LINHA DE GANHOS
    # =========================================================

    ganhos_row = resumo_row + 1

    worksheet.cell(
        row=ganhos_row,
        column=1,
        value="Ganhos",
    )

    worksheet.cell(
        row=ganhos_row,
        column=2,
        value=total_ganhos / 100,
    )

    # =========================================================
    # LINHA DE GASTOS
    # =========================================================

    gastos_row = resumo_row + 2

    worksheet.cell(
        row=gastos_row,
        column=1,
        value="Gastos",
    )

    worksheet.cell(
        row=gastos_row,
        column=2,
        value=total_gastos / 100,
    )

    # =========================================================
    # FLUXO TOTAL NA MESMA LINHA DOS GANHOS
    # =========================================================

    worksheet.cell(
        row=ganhos_row,
        column=4,
        value="FLUXO TOTAL",
    )

    worksheet.cell(
        row=ganhos_row,
        column=5,
        value=fluxo_total / 100,
    )

    worksheet.cell(
        row=ganhos_row,
        column=5,
    ).number_format = 'R$ #,##0.00'

    for column in (4, 5):
        cell = worksheet.cell(
            row=ganhos_row,
            column=column,
        )

        cell.font = Font(bold=True)
        cell.fill = total_fill
        cell.border = border

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    # =========================================================
    # LINHA DE LUCRO REAL
    # =========================================================

    lucro_row = resumo_row + 3

    worksheet.cell(
        row=lucro_row,
        column=1,
        value="Lucro real",
    )

    worksheet.cell(
        row=lucro_row,
        column=2,
        value=summary["lucro_real"] / 100,
    )

    # =========================================================
    # FORMATAÇÃO DO RESUMO
    # =========================================================

    for row in (
        ganhos_row,
        gastos_row,
        lucro_row,
    ):
        worksheet.cell(
            row=row,
            column=1,
        ).font = Font(bold=True)

        worksheet.cell(
            row=row,
            column=2,
        ).font = Font(bold=True)

        worksheet.cell(
            row=row,
            column=1,
        ).border = border

        worksheet.cell(
            row=row,
            column=2,
        ).border = border

        worksheet.cell(
            row=row,
            column=2,
        ).number_format = 'R$ #,##0.00'

    # Destaca o lucro real
    worksheet.cell(
        row=lucro_row,
        column=1,
    ).fill = total_fill

    worksheet.cell(
        row=lucro_row,
        column=2,
    ).fill = total_fill

    # =========================================================
    # FILTRO DAS TRANSAÇÕES
    # =========================================================
    #
    # IMPORTANTE:
    # Não usamos Table do Excel.
    # O AutoFilter simples evita o problema de reparo
    # que estava acontecendo no arquivo XLSX.

    if transactions:
        last_data_row = (
            first_data_row
            + len(transactions)
            - 1
        )

        worksheet.auto_filter.ref = (
            f"A{header_row}:G{last_data_row}"
        )

    # =========================================================
    # LARGURA DAS COLUNAS
    # =========================================================

    column_widths = {
        "A": 10,
        "B": 14,
        "C": 15,
        "D": 45,
        "E": 18,
        "F": 15,
        "G": 15,
    }

    for column, width in column_widths.items():
        worksheet.column_dimensions[column].width = width

    # =========================================================
    # CONGELAR CABEÇALHO
    # =========================================================

    worksheet.freeze_panes = "A5"

    # =========================================================
    # REMOVER GRADE DO EXCEL
    # =========================================================

    worksheet.sheet_view.showGridLines = False

    # =========================================================
    # SALVAR
    # =========================================================

    workbook.save(output_path)